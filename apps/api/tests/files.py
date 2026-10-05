"""Arquivos reais gerados em memória para os testes de ingestão."""

import io

import docx


def _pdf_escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[str]) -> bytes:
    """PDF mínimo e válido: uma página por texto, fonte Helvetica com acentuação (WinAnsi).

    Uma string vazia gera uma página sem texto, como a de um documento digitalizado.
    """
    page_ids = [4 + 2 * index for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    for page_id, text in zip(page_ids, pages, strict=True):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        operations = ["BT", "/F1 10 Tf", "13 TL", "40 800 Td"]
        operations += [f"({_pdf_escape(line)}) Tj T*" for line in text.split("\n") if line]
        operations.append("ET")
        stream = "\n".join(operations).encode("cp1252")
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))

    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
